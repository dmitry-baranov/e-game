"""Build editable PPTX, PDF, slide images and speaker notes from slides.json.

Run from the repository root:
  .venv-presentation/bin/python presentation/build_presentation.py
Requires python-pptx and Pillow. No external data or network calls.
"""
import json
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
SLIDES = json.loads((ROOT / "slides.json").read_text())
REPORT = json.loads((ROOT / "training-report-2026-10-09.json").read_text())
W, H, SCALE = 13.333333, 7.5, 120
BG = "F5F7FA"
INK = "142438"
MUTED = "526579"
TEAL = "087F8C"
TEAL_LIGHT = "E1F1F2"
WHITE = "FFFFFF"
LINE = "D9E2EC"
AMBER = "AA6520"
AMBER_LIGHT = "FFF2E2"
FONT_ROOT = Path("/System/Library/Fonts/Supplemental")
OVERFLOWS = []


def font(size, bold=False):
    name = "Arial Bold.ttf" if bold else "Arial.ttf"
    path = FONT_ROOT / name
    if not path.exists():
        path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold
                    else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    return ImageFont.truetype(str(path), round(size * SCALE / 72))


def wrap(text, width, size, bold=False):
    f = font(size, bold)
    lines = []
    for paragraph in text.split("\n"):
        line = ""
        for word in paragraph.split():
            candidate = (line + " " + word).strip()
            if f.getlength(candidate) <= width * SCALE or not line:
                line = candidate
            else:
                lines.append(line)
                line = word
        lines.append(line)
    return lines


class Canvas:
    def __init__(self, prs, index, dark=False):
        self.index = index
        self.slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = RGBColor.from_string(INK if dark else BG)
        self.image = Image.new("RGB", (1600, 900), "#" + (INK if dark else BG))
        self.draw = ImageDraw.Draw(self.image)
        self.dark = dark

    def rect(self, x, y, w, h, fill, outline=None, radius=False):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                           Inches(x), Inches(y), Inches(w), Inches(h))
        if radius:
            shape.adjustments[0] = 0.08
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        if outline:
            shape.line.color.rgb = RGBColor.from_string(outline)
        else:
            shape.line.fill.background()
        bbox = (round(x*SCALE), round(y*SCALE), round((x+w)*SCALE), round((y+h)*SCALE))
        if radius:
            self.draw.rounded_rectangle(bbox, radius=12, fill="#"+fill, outline="#"+outline if outline else None)
        else:
            self.draw.rectangle(bbox, fill="#"+fill, outline="#"+outline if outline else None)
        return shape

    def text(self, x, y, w, h, text, size=22, color=INK, bold=False, minimum=None):
        chosen = size
        minimum = minimum or size
        while True:
            lines = wrap(text, w, chosen, bold)
            line_h = chosen / 72 * 1.16
            # Additional descent space prevents clipping in PowerPoint.
            if len(lines)*line_h + chosen/72*0.10 <= h or chosen <= minimum:
                break
            chosen -= 0.5
        if len(lines)*line_h + chosen/72*0.10 > h + 0.02:
            OVERFLOWS.append((self.index, text, len(lines), chosen, h))
        tf = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "Arial"
            p.font.size = Pt(chosen)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = 1.16
            self.draw.text((x*SCALE, (y+i*line_h)*SCALE), line, font=font(chosen, bold), fill="#"+color, anchor="lt")
        return chosen

    def header(self, data):
        self.rect(.6, .55, .08, .25, TEAL)
        self.text(.85, .53, 11.8, .3, data["eyebrow"], 11.5, TEAL, True)
        self.text(.6, 1.05, 12.1, 1.08, data["title"], 31, INK, True, 28)

    def footer(self):
        color = "A9BACB" if self.dark else MUTED
        self.rect(.6, 7.04, 12.1, .012, "304258" if self.dark else LINE)
        self.text(.6, 7.15, 10.8, .24, "e-game · агенты рекомендаций · 09.10.2026", 10, color)
        self.text(12.05, 7.12, .65, .28, f"{self.index:02d} / {len(SLIDES)}", 10, color)

    def callout(self, text, amber=False):
        self.rect(.6, 6.18, 12.1, .63, AMBER_LIGHT if amber else TEAL_LIGHT, radius=True)
        self.text(.82, 6.32, 11.65, .4, text, 16.5, AMBER if amber else TEAL, True, 15)

    def card(self, x, y, w, h, title, body, title_size=21, body_size=19):
        self.rect(x, y, w, h, WHITE, LINE, radius=True)
        self.rect(x+.22, y+.24, .42, .05, TEAL)
        self.text(x+.22, y+.48, w-.44, .75, title, title_size, INK, True, 18)
        self.text(x+.22, y+1.26, w-.44, h-1.42, body, body_size, MUTED, minimum=17)

    def items(self, items, y=2.68, step=.75, size=23):
        for i, item in enumerate(items):
            self.rect(.68, y+i*step+.12, .07, .07, TEAL)
            self.text(.95, y+i*step, 11.7, step-.10, item, size, INK, minimum=19)

    def arrow(self, x1, y1, x2, y2, color=TEAL, dashed=False):
        connector = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
            Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        connector.line.color.rgb = RGBColor.from_string(color)
        connector.line.width = Pt(1.6)
        ln = connector._element.spPr.get_or_add_ln()
        if dashed:
            dash = OxmlElement('a:prstDash')
            dash.set('val', 'dash')
            ln.append(dash)
        end = OxmlElement('a:tailEnd')
        end.set('type', 'triangle')
        ln.append(end)
        a = (x1*SCALE, y1*SCALE)
        b = (x2*SCALE, y2*SCALE)
        if dashed:
            distance = math.dist(a,b)
            for offset in range(0, int(distance), 14):
                s, t = offset/distance, min(1,(offset+8)/distance)
                self.draw.line((a[0]+(b[0]-a[0])*s,a[1]+(b[1]-a[1])*s,
                                a[0]+(b[0]-a[0])*t,a[1]+(b[1]-a[1])*t), fill='#'+color,width=2)
        else:
            self.draw.line((a,b), fill='#'+color,width=2)
        angle = math.atan2(b[1]-a[1],b[0]-a[0])
        self.draw.polygon([b,(b[0]-10*math.cos(angle-.45),b[1]-10*math.sin(angle-.45)),
                          (b[0]-10*math.cos(angle+.45),b[1]-10*math.sin(angle+.45))], fill='#'+color)

    def diagram(self, data):
        for zone in data.get('zones', []):
            self.rect(*zone['box'], zone.get('fill','EAF0F5'), LINE, radius=True)
            self.text(zone['box'][0]+.16,zone['box'][1]+.12,zone['box'][2]-.32,.28,
                      zone['label'],11.5,MUTED,True)
        nodes = {n['id']:n for n in data['nodes']}
        for edge in data['edges']:
            if 'points' in edge:
                points = edge['points']
            else:
                a,b = nodes[edge['from']]['box'],nodes[edge['to']]['box']
                direction = edge.get('direction','right')
                if direction=='down': points=[(a[0]+a[2]/2,a[1]+a[3]),(b[0]+b[2]/2,b[1])]
                elif direction=='up': points=[(a[0]+a[2]/2,a[1]),(b[0]+b[2]/2,b[1]+b[3])]
                else: points=[(a[0]+a[2],a[1]+a[3]/2),(b[0],b[1]+b[3]/2)]
            for i in range(len(points)-1):
                self.arrow(*points[i],*points[i+1],edge.get('color',TEAL),edge.get('dashed',False))
            if edge.get('label'):
                x,y,w,h=edge['label_box']
                self.rect(x,y,w,h,BG)
                self.text(x,y,w,h,edge['label'],11.5,MUTED,minimum=10.5)
        for node in data['nodes']:
            x,y,w,h=node['box']
            tone=node.get('tone','default')
            fill=TEAL_LIGHT if tone=='ml' else (AMBER_LIGHT if tone=='future' else WHITE)
            color=TEAL if tone=='ml' else (AMBER if tone=='future' else INK)
            self.rect(x,y,w,h,fill,LINE,radius=True)
            self.text(x+.14,y+.12,w-.28,h-.20,node['text'],node.get('size',18),color,
                      node.get('bold',True),node.get('minimum',15.5))


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = "Агенты рекомендаций стратегий в e-game: шесть итераций"
    prs.core_properties.subject = "Результаты реализации, исследования и план ML + LLM + RAG"
    prs.core_properties.author = "Проект e-game"
    prs.core_properties.keywords = "e-game, агенты, стратегии, ML, LLM, RAG"
    pages = []
    previews = ROOT / "preview"
    previews.mkdir(exist_ok=True)
    md = ["# Агенты рекомендаций стратегий в e-game\n", f"{len(SLIDES)} слайдов. Ориентир выступления: 20–25 минут. Фактические данные на 09.10.2026.\n",
          "Строки «Текст слайда» — материал презентации; «Что говорить» — заметки докладчика.\n"]
    for index, data in enumerate(SLIDES, 1):
        layout = data["layout"]
        c = Canvas(prs, index, dark=layout in ("cover", "closing"))
        if layout == "cover":
            c.rect(.6, .65, .65, .065, "37BDC0")
            c.text(.6, 1.02, 11.8, .4, data["eyebrow"], 14, "77D4D5", True)
            c.text(.6, 1.91, 12.0, 2.12, data["title"], 46, WHITE, True)
            c.text(.6, 4.28, 11.9, .95, data["lead"], 25, "C4D5E5")
            c.text(.6, 5.73, 11.9, .72, "\n".join(data["items"]), 17, "A9BACB")
        elif layout == "closing":
            c.text(.6, .58, 12.0, .35, data["eyebrow"], 12, "77D4D5", True)
            c.text(.6, 1.1, 12.0, .7, data["title"], 34, WHITE, True)
            for i, item in enumerate(data["items"]):
                c.rect(.65, 2.17+i*.84, .07, .07, "77D4D5")
                c.text(.96, 2.04+i*.84, 11.7, .70, item, 22, WHITE, minimum=20)
            c.text(.6, 4.77, 12.0, 1.05, data["lead"], 24, "77D4D5", True, 22)
            c.text(.6, 6.3, 12.0, .45, data["callout"], 16, "C4D5E5")
        else:
            c.header(data)
            if layout == 'diagram':
                if data.get('lead'):
                    c.text(.6,2.02,12.1,.6,data['lead'],20,MUTED,minimum=18)
                c.diagram(data)
            elif layout == 'mockup':
                c.text(.6,2.01,12.1,.5,data['lead'],19.5,MUTED,minimum=18)
                for i,card in enumerate(data['cards']):
                    x=.6+i*4.1
                    c.rect(x,2.74,3.9,2.07,WHITE,LINE,radius=True)
                    c.text(x+.18,2.92,3.54,.4,card['title'],19,TEAL,True,17)
                    c.text(x+.18,3.47,3.54,1.18,card['body'],18,INK,minimum=16.5)
                c.rect(.6,5.04,12.1,.91,TEAL_LIGHT, radius=True)
                c.text(.8,5.19,11.7,.68,data['result'],19,INK,minimum=17.5)
            elif layout == 'example':
                c.text(.6,2.01,12.1,.57,data['lead'],19.5,MUTED,minimum=18)
                colw=data.get('widths',[3.6,2.2,2.2,4.1])
                rows=[data['columns']]+data['table']
                for row_i,row in enumerate(rows):
                    y=2.80+row_i*.58
                    c.rect(.6,y,12.1,.58,TEAL if row_i==0 else (WHITE if row_i%2 else 'EAF0F5'))
                    x=.6
                    for j,value in enumerate(row):
                        c.text(x+.16,y+.12,colw[j]-.28,.41,value,16.5,
                               WHITE if row_i==0 else INK,bold=row_i==0,minimum=15)
                        x+=colw[j]
                if data.get('result'):
                    c.text(.6,5.51,12.1,.49,data['result'],19,TEAL,True,minimum=17)
            elif layout in ("cards", "status"):
                cards = data["cards"]
                if data.get("lead"):
                    c.text(.6, 2.04, 12.1, .77, data["lead"], 22, MUTED, minimum=20)
                gap = .22
                width = (12.1-gap*(len(cards)-1))/len(cards)
                for i, card in enumerate(cards):
                    c.card(.6+i*(width+gap), 2.96, width, 2.86, card["title"], card["body"])
            elif layout == "timeline":
                for i, card in enumerate(data["cards"]):
                    x, y = .6+(i%3)*4.10, 2.38+(i//3)*1.74
                    c.rect(x, y, 3.9, 1.53, WHITE, LINE, radius=True)
                    c.text(x+.2, y+.2, 3.5, .48, card["title"], 18.5, TEAL, True, 17)
                    c.text(x+.2, y+.66, 3.5, .82, card["body"], 17, MUTED, minimum=15.5)
                c.text(.6, 6.28, 12.1, .42, "Итерация 6 поддерживает массовые исследования; технические детали вынесены за рамки доклада.", 15.5, MUTED)
            elif layout == "steps":
                for i, item in enumerate(data["items"]):
                    y = 2.36+i*.86
                    c.rect(.6, y+.04, .48, .48, TEAL, radius=True)
                    c.text(.755, y+.12, .27, .30, str(i+1), 17, WHITE, True)
                    c.text(1.30, y, 11.25, .76, item, 22, INK, minimum=19)
            elif layout == "pipeline":
                rows = data["rows"]
                start = 2.48 if len(rows) > 1 else 2.48
                for row_i, row in enumerate(rows):
                    y = start+row_i*1.48
                    c.text(.6, y, 12.1, .32, row["label"].upper(), 12, TEAL, True)
                    for i, step in enumerate(row["steps"]):
                        x = .6+i*3.12
                        c.rect(x, y+.48, 2.73, .92, WHITE, LINE, radius=True)
                        c.text(x+.17, y+.64, 2.39, .62, step, 20, INK, True, 18)
                        if i < 3:
                            c.text(x+2.81, y+.80, .29, .32, "→", 18, TEAL, True)
                if data.get("lead"):
                    lead_y = 2.04 if len(rows)>1 else 4.38
                    c.text(.6, lead_y, 12.1, .87, data["lead"], 22, MUTED, minimum=20)
            elif layout == "modelgrid":
                for i, card in enumerate(data["cards"]):
                    x, y = .6+(i%3)*4.10, 2.34+(i//3)*1.70
                    c.rect(x, y, 3.9, 1.5, WHITE, LINE, radius=True)
                    c.text(x+.20, y+.22, 3.5, .4, card["title"], 23, TEAL, True)
                    c.text(x+.20, y+.77, 3.5, .65, card["body"], 17, MUTED, minimum=16)
            elif layout == "series":
                c.text(.6, 2.04, 12.1, .5, data["lead"], 21, MUTED)
                colw = [5.3, 3.4, 3.4]
                rows = [data["columns"]]+data["table"]
                for row_i, row in enumerate(rows):
                    y = 2.75+row_i*.43
                    c.rect(.6, y, 12.1, .43, TEAL if row_i==0 else (WHITE if row_i%2 else "EAF0F5"))
                    x = .6
                    for j, value in enumerate(row):
                        c.text(x+.18, y+.085, colw[j]-.3, .34, value, 15.5,
                               WHITE if row_i==0 else INK, bold=row_i==0 or j>0)
                        x += colw[j]
            elif layout in ("stats", "split"):
                for i, stat in enumerate(data["stats"]):
                    x = .6+i*4.10
                    c.rect(x, 2.29, 3.9, 1.40, WHITE, LINE, radius=True)
                    c.text(x+.22, 2.46, 3.46, .70, stat["value"], 38, TEAL, True)
                    c.text(x+.22, 3.25, 3.46, .29, stat["label"], 15.5, MUTED)
                c.items(data["items"], y=4.0, step=.62, size=19.5)
            elif layout == "chart":
                c.text(.6, 2.02, 12.1, .43, data["lead"], 21, MUTED)
                names = ["LastSales", "Ridge", "CatBoost", "LightGBM", "RandomForest"]
                for i, name in enumerate(names):
                    val = REPORT["models"][name]["test_mae"]
                    y = 2.79+i*.59
                    c.text(.6, y+.04, 2.3, .34, name, 19, INK, name=="RandomForest")
                    length = val/18*8.0
                    c.rect(3.0, y, length, .40, TEAL if name=="RandomForest" else "90A8C0")
                    c.text(3.0+length+.14, y+.04, 1.0, .36, f"{val:.3f}".replace(".", ","), 18, TEAL if name=="RandomForest" else INK, True)
                c.text(3.0, 5.82, 9.3, .26, "Тест: 431 игра · 25 494 цикла · единицы товара", 13, MUTED)
            elif layout == "paired":
                c.text(.6, 2.04, 12.1, .78, data["lead"], 22, MUTED, minimum=20)
                for i, card in enumerate(data["cards"]):
                    c.card(.6+i*6.2, 3.00, 5.9, 2.08, card["title"], card["body"], body_size=19)
                c.text(.6, 5.39, 12.1, .45, "Δ = деньги после долгов модели − деньги после долгов правил", 21, TEAL, True, 19)
            elif layout == "roadmap":
                for i, card in enumerate(data["cards"]):
                    x, y = .6+(i%2)*6.20, 2.34+(i//2)*2.02
                    c.rect(x, y, 5.9, 1.82, WHITE, LINE, radius=True)
                    c.text(x+.22, y+.2, 5.46, .47, card["title"], 21, TEAL, True)
                    c.text(x+.22, y+.81, 5.46, .87, card["body"], 18, MUTED, minimum=17)
            else:
                raise ValueError(layout)
            if data.get("callout"):
                c.callout(data["callout"], amber=layout in ("paired", "status"))
        c.footer()
        source_text = "\n".join(data.get("sources", []))
        c.slide.notes_slide.notes_text_frame.text = data["notes"]+"\n\nИсточники в репозитории:\n"+source_text
        c.image.save(previews / f"slide-{index:02d}.png")
        pages.append(c.image)
        md += [f"\n## Слайд {index}. {data['title'].replace(chr(10), ' ')}\n", "### Текст слайда\n"]
        if data.get("lead"):
            md.append(data["lead"]+"\n")
        for card in data.get("cards", []):
            md.append(f"- **{card['title']}** — {card['body']}")
        for item in data.get("items", []):
            md.append("- "+item)
        for stat in data.get("stats", []):
            md.append(f"- **{stat['value']}** — {stat['label']}")
        for row in data.get("rows", []):
            md.append(f"- {row['label']}: "+" → ".join(s.replace("\n", " ") for s in row["steps"]))
        for node in data.get('nodes', []):
            md.append(f"- **{node['id']}**: {node['text'].replace(chr(10),' ')}")
        for edge in data.get('edges', []):
            if 'from' in edge:
                md.append(f"- Связь: {edge['from']} → {edge['to']}"+(f" ({edge['label']})" if edge.get('label') else ''))
        if data.get('result'):
            md.append('\n**Результат:** '+data['result'])
        if data.get("table"):
            md += ["", "| "+" | ".join(data["columns"])+" |", "|"+'|'.join(['---']*len(data['columns']))+"|"]
            md += ["| "+" | ".join(row)+" |" for row in data["table"]]
        if layout=="chart":
            md += ["", "| Модель | MAE validation | MAE test |", "|---|---:|---:|"]
            for name in names:
                m = REPORT["models"][name]
                md.append(f"| {name} | {m['validation_mae']:.3f} | {m['test_mae']:.3f} |")
        if data.get("callout"):
            md.append("\n**Акцент:** "+data["callout"])
        md += ["\n### Что говорить\n", data["notes"], "\n### Основания\n"]
        for path in data["sources"]:
            md.append(f"- [{path}](../{path})")
    if OVERFLOWS:
        raise RuntimeError("Text overflow: "+repr(OVERFLOWS))
    output = ROOT / "e-game_strategy_agents_6_iterations"
    prs.save(str(output.with_suffix(".pptx")))
    pages[0].save(output.with_suffix(".pdf"), save_all=True, append_images=pages[1:], resolution=SCALE)
    (ROOT / "speaker-notes.md").write_text("\n".join(md)+"\n")
    overview = Image.new("RGB", (1600, math.ceil(len(pages)/5)*200), "#E3E9EF")
    draw = ImageDraw.Draw(overview)
    for i, page in enumerate(pages):
        thumb = page.resize((312, 176), Image.Resampling.LANCZOS)
        x, y = (i%5)*320+4, (i//5)*200+4
        overview.paste(thumb, (x, y))
        draw.text((x+4, y+181), f"{i+1:02d}", fill="#142438", font=font(8))
    overview.save(ROOT / "overview.png")
    # Verify the generated archive opens, notes exist and all sources resolve.
    reopened = Presentation(str(output.with_suffix(".pptx")))
    assert len(reopened.slides) == len(SLIDES)
    assert all(s.notes_slide.notes_text_frame.text.strip() for s in reopened.slides)
    assert REPORT["status"] == "trained" and REPORT["games"] == 2152
    for data in SLIDES:
        for source in data["sources"]:
            assert (ROOT.parent / source).exists(), source
    print(f"Built {len(pages)} slides: editable PPTX, PDF, previews, speaker notes. No text overflows.")


if __name__ == "__main__":
    build()
